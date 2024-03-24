import streamlit as st
from pptx import Presentation
from PyPDF2 import PdfReader
from docx import Document
import os
from langchain_google_genai import ChatGoogleGenerativeAI
from langchain.chains.question_answering import load_qa_chain
from dotenv import load_dotenv

load_dotenv()
os.getenv("GOOGLE_API_KEY")
genai.configure(api_key=os.getenv("GOOGLE_API_KEY"))


def get_file_text(files):
    text = ""
    for file in files:
        file_extension = file.name.split(".")[-1].lower()
        if file_extension == "pdf":
            for page in pdf_reader.pages:
                text += page.extract_text()
        elif file_extension == "pptx":
            prs = Presentation(file)
                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        text += shape.text
        elif file_extension == "docx":
            for paragraph in doc.paragraphs:
                text += paragraph.text
        else:
            st.error(f"Unsupported file format: {file.name}")
    return text


def get_text_chunks(text):
    text_splitter = RecursiveCharacterTextSplitter(chunk_size=10000, chunk_overlap=1000)
    chunks = text_splitter.split_text(text)
    return chunks


def get_vector_store_safe(text_chunks):
    embeddings = GoogleGenerativeAIEmbeddings(model="models/embedding-001")
    vector_store = FAISS.from_texts(text_chunks, embedding=embeddings)
    vector_store.save_local("faiss_index_safe")


def get_conversational_chain():
    prompt_template = """
    Answer the question as detailed as possible from the provided context, make sure to provide all the details, if the answer is not in
    provided context just say, "answer is not available in the context", don't provide the wrong answer\n\n
    Context:\n {context}?\n
    Question: \n{question}\n

    Answer:
    """

    model = ChatGoogleGenerativeAI(model="gemini-pro", temperature=0.3)

    return chain


def main():
    st.set_page_config("CONVERSADOCS")
    st.header("HOLA USER!!!")
    st.header("Chat with Documents using CONVERSADOCS💁")

    if st.sidebar.button("Submit"):
        if not files:
            st.sidebar.warning("Please upload at least one file.")
        else:
            with processing_status:
                status_text = "Processing..."
                text = get_file_text(files)
                status_text = "Done"

    user_question = st.text_input("Ask a Question from the Documents")

    if st.button("Ask Question"):
        if not user_question:
            st.warning("Please enter a question.")
        else:
            with st.spinner("Processing..."):
                embeddings = GoogleGenerativeAIEmbeddings(model="models/embedding-001")
                new_db = get_faiss_index_safe(embeddings)
                response = chain(
                    {"input_documents": docs, "question": user_question},
                    return_only_outputs=True
                )
                st.write("Reply: ", response["output_text"])


if __name__ == "__main__":
    main()
